"""Báo cáo Customer theo mẫu TV3, chỉ sử dụng bằng chứng Customer đã chạy."""
from pathlib import Path
from copy import deepcopy
import sys, json, hashlib, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
REFERENCE=Path(sys.argv[1])
OUT=ROOT/'docs/BAO_CAO_TONG_HOP_TV2_CUSTOMER.docx'
source_hash=hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
# Load verified task facts without executing the older report's output commands.
facts={'__file__':str(ROOT/'scripts/build-customer-report.py')}
source=(ROOT/'scripts/build-customer-report.py').read_text(encoding='utf-8')
exec(compile(source.split('# Các tài liệu Markdown')[0],str(ROOT/'scripts/build-customer-report.py'),'exec'),facts)
tasks=facts['tasks']; checklist=facts['checklist']; E=facts['EVIDENCE']
tasks[-1]['pass']='Test và tài liệu local đã có; một commit core local f36a1bc. Chưa push, chưa mở PR và chưa xác nhận CI/review.'
tasks[-1]['checks'][-1]=('Bàn giao Git','Commit có danh tính; nhánh remote/PR/CI được xác minh riêng','Có commit core local f36a1bc; phần còn lại chờ Hưng xem báo cáo trước')
checklist[-1]=(checklist[-1][0],False,'Đã có báo cáo và ảnh thật; chưa có PR, CI hoặc lead review.')
d=Document(REFERENCE)
cover=[deepcopy(p._p) for p in d.paragraphs[:8]]
prototypes={len(t.columns):deepcopy(t._tbl) for t in d.tables[:11]}
meta=deepcopy(d.tables[0]._tbl)
sect=deepcopy(d.sections[0]._sectPr)
for child in list(d._element.body):d._element.body.remove(child)
d._element.body.append(sect)
for rid,rel in list(d.part.rels.items()):
    if rel.reltype.endswith('/image'):d.part.drop_rel(rid)
def add_element(el): d._element.body.insert(len(d._element.body)-1,el)
def p(text='',style=None): return d.add_paragraph(text,style)
def h(text,level=1): return d.add_heading(text,level)
def page(): d.add_page_break()
def table(headers,rows,widths):
    el=deepcopy(prototypes[len(headers)]);add_element(el);t=Table(el,d._body)
    for row in list(t.rows)[2:]:el.remove(row._tr)
    row_proto=deepcopy(t.rows[1]._tr)
    while len(t.rows)<len(rows)+1:el.append(deepcopy(row_proto))
    for ri,values in enumerate([headers]+rows):
        row=t.rows[ri];pr=row._tr.get_or_add_trPr()
        for height in list(pr.findall(qn('w:trHeight'))):pr.remove(height)
        if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
        for ci,value in enumerate(values):
            cell=row.cells[ci];cell.width=Inches(widths[ci])
            cp=deepcopy(cell.paragraphs[0]._p.pPr)
            rp=deepcopy(cell.paragraphs[0].runs[0]._r.rPr) if cell.paragraphs[0].runs else None
            cell.text=str(value);para=cell.paragraphs[0]
            if cp is not None:
                if para._p.pPr is not None:para._p.remove(para._p.pPr)
                para._p.insert(0,cp)
            para.paragraph_format.keep_with_next=False
            para.paragraph_format.space_after=Pt(3)
            para.paragraph_format.space_before=Pt(3)
            para.paragraph_format.line_spacing=1.0
            para.alignment=WD_ALIGN_PARAGRAPH.LEFT
            if rp is not None:para.runs[0]._r.insert(0,rp)
            para.runs[0].font.size=Pt(10)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    p().paragraph_format.space_after=Pt(0)
    return t
for i,el in enumerate(cover):
    for node in el.iter(qn('w:t')):
        node.text=node.text.replace('THÀNH VIÊN 3 — ADMIN OPERATIONS','THÀNH VIÊN 2 — CUSTOMER').replace('ADM-01 → ADM-07','CUS-01 → CUS-07')
    add_element(el)
    if i==3:
        d.paragraphs[-1].style=d.styles['Title']
        d.paragraphs[-1].paragraph_format.space_before=Pt(0)
        d.paragraphs[-1].paragraph_format.space_after=Pt(6)
        d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        for run in d.paragraphs[-1].runs:run.font.name='Times New Roman'
        for border in list(d.styles['Title'].element.iter(qn('w:pBdr'))):border.getparent().remove(border)
add_element(meta);mt=Table(meta,d._body)
metadata=[('Người thực hiện','Thành viên 2: Hưng · MSSV: chưa cung cấp'),('Repo','github.com/tbao4151/Lap_trinh_web'),('Nhánh gốc','feature/requirement-2-php-mysql · 9c35a94'),('Nhánh làm việc','feature/customer-completion · chưa push'),('Ngày lập','15/10/2026')]
for row,values in zip(mt.rows,metadata):
    for cell,value in zip(row.cells,values):cell.text=value
p('Bản trình Hưng xem trước khi tiếp tục đưa nhánh lên GitHub.')
page()
toc_heading=p('MỤC LỤC');toc_heading.alignment=WD_ALIGN_PARAGRAPH.CENTER;toc_heading.runs[0].bold=True
toc=p();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');toc._p.append(field)
page();h('1. Tóm tắt kết quả')
p('Báo cáo đối chiếu công việc Customer của thành viên 2 với 02_PLAN_THANH_VIEN_2_CUSTOMER.md. Phạm vi gồm CUS-01 đến CUS-07, tổng 27,5 điểm kế hoạch. Bản này dùng cấu trúc trình bày của báo cáo TV3; các số liệu, ảnh và kết quả bên dưới đều thuộc phần Customer.')
for s in ['18 trang khách hàng được kiểm tra HTTP và 5 kích thước màn hình; 90/90 lượt responsive đạt.', 'HTTP và SQL nghiệp vụ 35/35; browser 103/103; PHP lint 50/50; JavaScript 8/8; SQL tổng hợp 9/9; kiểm thử bổ sung 10/10.', 'Đã sửa thông báo khoảng giá và ngân sách, nhãn sắp xếp, trạng thái bận khi lưu, giữ dữ liệu đăng ký hợp lệ và tương phản nút.', '5 task đạt LOCAL_PASS, 2 task BLOCKED. Tổng trọng số đạt local là 19,5/27,5; không phải điểm giảng viên đã chấm.']:
    p(s,'List Bullet')
p('CUS-02 còn gallery nhiều ảnh và contract tồn kho. CUS-07 còn nhánh remote, PR, CI và nghiệm thu. Một commit core đã tạo local trước khi Hưng yêu cầu xem Word trước; chưa đưa thay đổi lên GitHub.')
h('1.1. Bảng đối chiếu với kế hoạch',2)
table(['Hạng mục','Điểm KH','Trạng thái','Minh chứng chính','Tự đánh giá'],[[t['id']+' '+t['title'],t['points'],t['status'],'Mục 3.'+str(i+1)+'; evidence/'+t['id'],t['points'] if t['status']=='LOCAL_PASS' else 'Chưa tính'] for i,t in enumerate(tasks)]+[['Tổng','27,5','5 local đạt','2 task còn điều kiện','19,5 / 27,5']],[1.65,.55,.9,2.35,.85])
page();h('2. Phạm vi môi trường và nguyên tắc làm việc')
table(['Mục','Chi tiết'],[['Mã nguồn','Nhánh Customer từ feature/requirement-2-php-mysql tại 9c35a94. Không triển khai chức năng Admin theo mẫu TV3.'],['PHP và CSDL','PHP 8.2.12 XAMPP; MariaDB cổng 3307; cas_hoa_customer_preview riêng.'],['Ứng dụng','PHP development server http://127.0.0.1:8033; Chrome qua Playwright; SMTP thử cổng 1025.'],['Dữ liệu','Tài khoản, địa chỉ và đơn kiểm thử; thay đổi tạm sản phẩm được khôi phục.'],['Phạm vi chưa xác minh','Apache tại URL môn học, SMTP Internet, CI và nghiệm thu sau tích hợp.']],[1.3,5])
p('Các kết quả được lấy từ log thực chạy ngày 14–15/10/2026. Ảnh chụp nội dung trang và vùng giao diện thật, không chứa thanh địa chỉ Chrome. Không dùng ảnh Admin, số liệu của thành viên 3 hoặc dựng thanh địa chỉ giả.')
p('Các diff checkout/auth đã tồn tại trước lượt hoàn thiện cuối và giữ nguyên SHA256 trong lượt đó. Commit f36a1bc ghi nhận riêng phần core để lead review. Không thay schema, role, trạng thái đơn hay luật trừ kho trong lượt làm báo cáo.')
page();h('3. Chi tiết từng hạng mục')
for i,t in enumerate(tasks):
    if i:page()
    h(f"3.{i+1}. {t['id']} {t['title']}",2)
    p('Trọng số kế hoạch: '+t['points']+' điểm. Trạng thái: '+t['status']+'.')
    p('Yêu cầu kế hoạch: '+t['goal'])
    p('Đã thực hiện: '+t['actual'])
    table(['Kiểm tra','Kỳ vọng','Kết quả thực tế'],t['checks'],[1.25,2.35,2.7])
    p('Đối chiếu SQL: '+t['sql_note'])
    p('Tệp liên quan: '+', '.join(t['files'])+'.')
    p('Điều kiện nghiệm thu: '+t['pass'])
    p('Ảnh minh chứng: xem Phụ lục A và docs/evidence/'+t['id']+'/ trong repo.')
    p('Commit đề xuất theo task: '+t['commit']+'.')
page();h('4. Đối chiếu checklist nghiệm thu')
table(['STT','Yêu cầu','Trạng thái','Minh chứng'],[[i+1,label,'Đạt local' if done else 'Còn thiếu',proof] for i,(label,done,proof) in enumerate(checklist)],[.35,2.55,.8,2.6])
h('5. Đối chiếu các giới hạn công việc')
table(['Nguyên tắc','Thực hiện và giới hạn'],[['Không thay schema hoặc luật core tùy ý','Không có migration mới. Core có sẵn được giữ nguyên trong lượt cuối và tách commit để review.'],['Không giả kết quả','Số liệu lấy từ JSON; giữ lại lần lỗi harness, timeout và dịch vụ dừng.'],['Không công khai thông tin nhạy cảm','Không đưa .env, hộp thư reset, token hay dữ liệu database runtime vào báo cáo.'],['Không làm thay Admin','Tài liệu TV3 chỉ là mẫu; không sao chép kết quả 216 test hoặc 33 ảnh Admin.'],['Git theo yêu cầu mới nhất','Đã có một commit local; dừng thao tác Git để trình Word trước. Chưa push hoặc mở PR.']],[2.2,4.1])
page();h('6. Kết quả kiểm thử')
table(['Bộ kiểm tra','Đạt','Nguồn'],[[label,facts['count'](facts[key]),path] for label,key,path in [('HTTP và SQL nghiệp vụ','http','evidence/customer/http-results.json'),('Browser và responsive','browser','evidence/customer/browser-results.json'),('PHP lint','php','evidence/CUS-07/php-lint.json'),('JavaScript syntax','js','evidence/CUS-07/js-lint.json'),('SQL tổng hợp','sql','evidence/CUS-07/sql-results.json')]]+[['Bổ sung theo task','10/10','evidence/CUS-02..CUS-06/results.json']],[1.7,.7,3.9])
p('Các bộ test có thể kiểm tra cùng một nghiệp vụ ở nhiều tầng; không cộng tổng số assertion thành số chức năng độc lập. Không dùng CI cũ của nhánh gốc để tuyên bố CI của thay đổi Customer đã đạt.')
p('Lỗi đã gặp: khai báo crypto lặp trong test HTTP; điều hướng fragment không phát sinh response; timeout 7 giây ở trang tài khoản; preview dừng sau gián đoạn phiên. Đã sửa harness hoặc khởi động lại đúng database hiện có rồi chạy lại đạt. Bằng chứng thất bại vẫn được giữ trong CUS-07.')
h('7. Thay đổi mã nguồn')
table(['Nhóm tệp','Vai trò','Phạm vi'],[['includes/customer.php, customer-recommendations.php','Helper Customer và gợi ý DB','Tệp mới của phần Customer'],['assets/css/customer.css; customer-catalog.js; customer-forms.js','Giao diện, bộ lọc, validation','Tệp mới; có sửa bổ sung lượt cuối'],['18 trang PHP khách hàng và assets/js/ liên quan','Nối PHP/Session/DB, trải nghiệm và phản hồi','Gồm thay đổi có sẵn và lượt hoàn thiện'],['actions/cart.php; actions/checkout.php; includes/auth.php; bootstrap.php','Validation, wrapping, replay, snapshot','Core có sẵn; commit f36a1bc để lead review'],['tests/customer/; scripts/customer-*','HTTP/browser/SQL, môi trường thử riêng','Không ghi dữ liệu thật'],['docs/ và scripts/build-customer-report*.py','Audit, test, Word, minh chứng','Tài liệu và công cụ dựng báo cáo']],[2.4,2.3,1.6])
p('Không nhận toàn bộ diff trong workspace là code mới của lượt cuối. Phần sửa bổ sung cụ thể đã nêu tại mỗi CUS; các thay đổi trước đó được bảo toàn, không ghi đè.')
page();h('8. Việc còn lại và đề xuất')
table(['Mã','Việc cần xử lý','Điều kiện'],[['CUS-02','Gallery nhiều ảnh và trạng thái hết hàng thành phẩm','Lead chốt schema/quan hệ định mức; tích hợp rồi kiểm thử lại'],['CORE','Review wrapping, snapshot, chống đơn trùng và validation','Xem commit local f36a1bc; chưa coi đã được lead duyệt'],['CUS-07','Commit phần còn lại, đưa nhánh lên GitHub, PR và CI','Hưng xem Word trước theo yêu cầu mới nhất'],['ENV','Apache và email Internet','Kiểm tra trên môi trường nhóm; local không thay thế nghiệm thu']],[.75,3,2.55])
h('8.1. Trạng thái Git tại thời điểm lập báo cáo',2)
p('Nhánh local: feature/customer-completion. Commit local f36a1bc — fix(customer): preserve validated cart and checkout session behavior. Tác giả dùng tên Hưng và email do Hưng cung cấp. Chưa push, chưa tạo nhánh remote Customer, chưa có PR. Phần còn lại chưa commit.')
p('Kế hoạch tiếp theo: sau khi trình bản Word, tiếp tục các thao tác Git theo chỉ dẫn của Hưng; nhánh tích hợp dự kiến feature/requirement-2-php-mysql. Không merge vào main trong bước bàn giao này.')
h('8.2. Đánh giá đóng góp',2)
p('19,5/27,5 là tổng trọng số của năm task có đủ bằng chứng local. Hai task chưa đủ điều kiện được giữ BLOCKED. Tỷ lệ này không phải điểm đã chấm, số dòng mã hay tỷ lệ GitHub Contributors; lead cần review và nghiệm thu tích hợp.')
page();h('Phụ lục A. Ảnh minh chứng Customer')
p('Ảnh từ Chrome do Playwright điều khiển trong môi trường thử PHP/MySQL. Ảnh nội dung trang hoặc vùng UI, không có thanh địa chỉ. Số liệu và thông tin hiển thị là dữ liệu kiểm thử. Mỗi hình ghi nguồn để đối chiếu ảnh gốc.')
figures=[('CUS-01/report-home.png','Trang chủ trên màn hình 1024 px'),('CUS-02/report-catalog-invalid-price.png','Thông báo khi khoảng giá bị đảo'),('CUS-03/report-checkout-retry.png','Checkout báo lỗi thử nghiệm và cho phép thử lại'),('CUS-04/report-reset-success.png','Khôi phục mật khẩu thành công qua SMTP local'),('CUS-05/report-login-contrast.png','Đăng nhập với màu nút đã đạt tương phản 6,32:1'),('CUS-06/report-recommendation-empty.png','Gợi ý không có kết quả và nút đặt lại bộ lọc'),('CUS-07/report-order-success.png','Đơn được tạo thật và có mã tra cứu'),('CUS-04/login-error-375.png','Đăng nhập sai có thông báo trên màn hình 375 px'),('CUS-02/catalog-search-1440.png','Kết quả tìm kiếm và bộ lọc trên desktop'),('CUS-03/cart-1440.png','Giỏ hàng đã đồng bộ Session sau tải lại'),('CUS-04/profile-1440.png','Hồ sơ tài khoản Customer'),('customer/catalog-filter-375.png','Bộ lọc sản phẩm trên mobile'),('customer/lookup-375.png','Tra cứu đơn theo mã và số điện thoại')]
for i,(rel,label) in enumerate(figures):
    if i and i not in (2,6):page()
    image_path=E/rel
    with Image.open(image_path) as im:w,hh=im.size
    width=min(6.05,7.5*w/hh)
    if i==1:width=2.8
    ip=p();ip.alignment=WD_ALIGN_PARAGRAPH.CENTER;ip.paragraph_format.keep_with_next=True
    ip.add_run().add_picture(str(image_path),width=Inches(width))
    cap=p(f'Hình {i+1}. {label}','Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic=True;cap.runs[0].bold=False;cap.runs[0].font.size=Pt(11);cap.runs[0].font.color.rgb=RGBColor(0,0,0)
    src=p('Nguồn: docs/evidence/'+rel);src.runs[0].font.size=Pt(9)
page();h('Phụ lục B. Danh sách kiểm tra chức năng')
p('Các dòng dưới lấy nguyên tên kiểm tra từ JSON kết quả. Ma trận responsive 90 lượt được tóm tắt ở phần audit; không lặp toàn bộ 90 dòng tại đây.')
rows=[]
for label,group in [('HTTP',facts['http']),('Browser',facts['browser'])]+list(facts['supplement'].items()):
    for r in group['results']:
        if r['name'].startswith('Responsive '):continue
        rows.append([label,r['name'],r['status']])
table(['Nhóm','Kiểm tra','Kết quả'],rows,[.7,4.9,.7])
page();h('Audit 18 trang khách hàng',2)
audit=[]
for r in facts['http']['results']:
    if r['name'].startswith('HTTP page '):audit.append([r['name'][10:].split('?')[0],'LOCAL_PASS','HTTP 200; 5 viewport; ảnh 375/768/1440'])
table(['Trang PHP','Trạng thái','Bằng chứng'],audit,[2.45,1.05,2.8])
page();h('Phụ lục C. SQL kiểm tra và kết quả')
p('Chỉ xuất truy vấn tổng hợp và số đếm; không xuất mật khẩu, token hoặc thông tin cá nhân. Đơn TEST-* dùng cho phân trang không được cộng vào kiểm tra tổng tiền đơn CSH-*.')
for r in facts['sql']['results']:
    para=p(r['name']);para.runs[0].bold=True;para.paragraph_format.keep_with_next=True
    query=p(r['sql']+';');query.runs[0].font.name='Consolas';query.runs[0].font.size=Pt(9)
    p('Kỳ vọng: '+str(r.get('expected','>= '+str(r.get('minimum'))))+'. Thực tế: '+str(r['actual'])+'. '+r['status']+'.')
d.core_properties.title='Báo cáo tổng hợp công việc thành viên 2 Customer'
d.core_properties.author='Hưng'
d.core_properties.last_modified_by='Hưng'
d.save(OUT)
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest()==source_hash
assert not any('ADM-' in p.text or 'THÀNH VIÊN 3' in p.text for p in d.paragraphs)
assert len(d.inline_shapes)==len(figures)
(ROOT/'.local/tv3-template/build.json').write_text(json.dumps({'reference_sha256':source_hash,'images':len(figures),'tables':len(d.tables),'output':str(OUT)},ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'output':str(OUT),'images':len(figures),'tables':len(d.tables)},ensure_ascii=False))
