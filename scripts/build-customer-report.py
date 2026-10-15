"""Tạo tài liệu bàn giao tiếng Việt từ bằng chứng đã chạy, không suy diễn kết quả."""
from pathlib import Path
import json
import sys
sys.stdout.reconfigure(encoding='utf-8')
from datetime import datetime, timezone, timedelta
from collections import Counter
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
EVIDENCE = DOCS / 'evidence'
def read_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))
http = read_json(EVIDENCE/'customer/http-results.json')
browser = read_json(EVIDENCE/'customer/browser-results.json')
php = read_json(EVIDENCE/'CUS-07/php-lint.json')
js = read_json(EVIDENCE/'CUS-07/js-lint.json')
sql = read_json(EVIDENCE/'CUS-07/sql-results.json')
supplement = {f'CUS-0{i}':read_json(EVIDENCE/f'CUS-0{i}/results.json') for i in range(2,7)}
for group in [http,browser,php,js,sql,*supplement.values()]:
    assert all(r['status']=='PASS' for r in group['results']), 'Còn test chưa đạt; không tự xuất kết luận pass'
guards=read_json(EVIDENCE/'CUS-07/protected-files.json')
assert all(r['status']=='PASS' for r in guards)
tasks = [
 {'id':'CUS-01','title':'Audit 18 trang khách hàng','points':'3','status':'LOCAL_PASS',
  'goal':'Kiểm tra HTTP, PHP, JS, CSS, ảnh, form/SQL và link PHP trên 18 trang.',
  'actual':'Audit đủ 18 trang, 90 lượt trang/viewport; ảnh thật ở 375, 768, 1440. Không thấy PHP warning, lỗi JS chưa bắt, ảnh hỏng hay link .html trong phạm vi kiểm thử.',
  'files':['docs/customer-audit.md','PROGRESS.md','docs/evidence/CUS-01/'],
  'pass':'Đạt smoke và audit local; Apache URL yêu cầu chưa đạt.',
  'commit':'docs(customer): audit 18 customer pages with local evidence',
  'image':'CUS-01/report-home.png',
  'checks':[('18 trang được quyền','HTTP 200, không đổi sang trang đăng nhập ngoài ý muốn','18/18 đạt; slug sản phẩm sai trả 404'),('5 kích thước','375/390/768/1024/1440 không tràn ngang, có nhãn và ảnh tải','90/90 lượt đạt trong bộ browser cuối')],
  'sql_note':'Dùng kiểm tra dữ liệu nền, sản phẩm active và bài công khai; form liên hệ và tài khoản được đối chiếu bằng các test HTTP/SQL.'},
 {'id':'CUS-02','title':'Catalog và chi tiết sản phẩm','points':'5','status':'BLOCKED',
  'goal':'Lọc DB đa tiêu chí, sắp xếp, cập nhật giá/ẩn; detail, ảnh, gói hoa và trạng thái hết hàng.',
  'actual':'Tái sử dụng catalog DB và giấy gói có sẵn. Sửa nhãn Mới thêm; thêm lỗi giá inline và aria-invalid. Đối chiếu 6 bộ lọc riêng với tập ID SQL, kiểm tra tổ hợp, giá mới và hidden.',
  'files':['san-pham.php','assets/js/customer-catalog.js','tests/customer/completion.cjs','docs/de-xuat-cho-lead.md'],
  'pass':'Catalog hiện có đạt local; chưa nghiệm thu đủ gallery/tồn kho thành phẩm vì cần lead chốt schema/contract.',
  'commit':'fix(customer): clarify catalog sorting and price validation',
  'image':'CUS-02/report-catalog-invalid-price.png',
  'checks':[('Bộ lọc riêng và kết hợp','Tập sản phẩm đúng như SQL, URL giữ được','6 bộ lọc riêng khớp SQL; bộ lọc kết hợp và tìm không dấu đạt'),('Đổi giá và ẩn','Giá mới hiện trên detail; hidden không có trong catalog, slug 404','Giá thử 456000 đã hiện; hidden bị loại; khôi phục giá/trạng thái cũ'),('Khoảng giá đảo','Hiện lỗi rõ và đặt lại xóa lỗi','Lỗi inline/aria-invalid đạt; nhãn sắp xếp Mới thêm đúng dữ liệu')],
  'sql_note':'SELECT sản phẩm theo type/category/tone/occasion/price/slug. UPDATE tạm giá/status trong DB thử, khôi phục trong finally. Chưa kiểm thử thao tác bấm nút Admin; đã xác minh tác động dữ liệu Admin sử dụng.'},
 {'id':'CUS-03','title':'Giỏ hàng và trải nghiệm checkout','points':'5','status':'LOCAL_PASS',
  'goal':'Session giỏ bền qua refresh, lỗi mạng/retry, giao hàng/pickup, giá server, chống đơn trùng.',
  'actual':'Giữ xử lý Session/checkout có sẵn, bổ sung aria-busy khi lưu và xóa trạng thái bận khi lỗi. Kiểm thử giỏ/checkout với dữ liệu thật trong DB thử.',
  'files':['assets/js/main-r2.js','tests/customer/completion.cjs','docs/evidence/CUS-03/'],
  'pass':'Đạt trên bản local đang có. Cơ chế idempotency và snapshot trong checkout là diff từ trước phiên, cần lead review riêng.',
  'commit':'fix(customer): improve cart and checkout feedback',
  'image':'CUS-03/report-checkout-retry.png',
  'checks':[('Giỏ và lỗi mạng','Refresh giữ dữ liệu, cập nhật/xóa đúng, lỗi không giả thành công','Giỏ 1 dòng số lượng 3 sau refresh; retry và mất feed đạt'),('Giá và đơn trùng','Không tin giá gửi từ trình duyệt; hai request cùng key chỉ 1 đơn','2 mẫu p001 có tổng 800000; hai response cùng mã và DB chỉ thêm 1 đơn'),('Ship và pickup','Địa chỉ chỉ bắt buộc khi giao; ngày/giờ/phone được kiểm tra','Giao hàng và pickup đều lưu; giá trị sai bị 422; lỗi 503 inline và có thể thử lại')],
  'sql_note':'SELECT orders/order_items theo mã thật, kiểm số đơn tăng đúng 1; kiểm tổng tiền, số dòng, phương thức giao nhận. Không sửa file actions/checkout.php trong phiên.'},
 {'id':'CUS-04','title':'Tài khoản hồ sơ và khôi phục','points':'5','status':'LOCAL_PASS',
  'goal':'Validation, đăng nhập/đăng xuất, hồ sơ, lịch sử phân trang, cách ly A/B, lookup và reset email.',
  'actual':'Đăng ký giữ tên/email/phone khi server từ chối và vẫn để trống mật khẩu. Bổ sung ràng buộc UI; kiểm tra reset qua email SMTP local và trình duyệt.',
  'files':['dang-ky.php','tests/customer/completion.cjs','docs/evidence/CUS-04/'],
  'pass':'Đạt SMTP local, token một lần/hết hạn và cách ly tài khoản. Chưa kiểm chứng thư gửi tới nhà cung cấp Internet.',
  'commit':'fix(customer): retain signup details and verify account recovery',
  'image':'CUS-04/report-reset-success.png',
  'checks':[('Validation và tài khoản','Dữ liệu sai không tạo user; email trùng báo rõ; profile lưu DB','4 loại input sai bị từ chối; số user không đổi; đăng ký/login/logout/profile đạt'),('Lịch sử và lookup','A không xem được đơn B; phân trang; phone sai không lộ đơn','10 đơn/trang; trang 2 có đơn cũ; A/B cách ly; lookup phone sai 404'),('Reset qua email','SMTP nhận thư, token không hiện URL/UI, dùng một lần và hết hạn','Thư local nhận thật; reset qua browser đạt; replay/hết hạn bị từ chối; SQL consumed_at có giá trị')],
  'sql_note':'SELECT users/customers sau profile, JOIN orders theo user để kiểm liên kết; password_resets chỉ kiểm số bản ghi/consumed_at. Không ghi mật khẩu hoặc token vào tài liệu.'},
 {'id':'CUS-05','title':'Responsive và khả năng truy cập','points':'4','status':'LOCAL_PASS',
  'goal':'5 viewport; menu, filter, form, giỏ, empty/error/success; label, focus, contrast và dữ liệu dài.',
  'actual':'Giữ bố cục hiện có, chỉnh token màu riêng cho Customer. Kiểm bàn phím mở bộ lọc, focus visible, ảnh và tràn ngang.',
  'files':['assets/css/customer.css','tests/customer/browser.cjs','tests/customer/completion.cjs','docs/evidence/CUS-05/'],
  'pass':'Đạt các kiểm tra cụ thể đã chạy, không tuyên bố chứng nhận WCAG toàn site hoặc kiểm thử screen reader toàn diện.',
  'commit':'fix(customer): improve contrast and keyboard accessibility',
  'image':'CUS-05/report-login-contrast.png',
  'checks':[('Tương phản nút','Tỷ lệ chữ nút >=4,5:1','Trước sửa 3,94:1; sau sửa 6,32:1'),('Responsive và bàn phím','Không overflow, label/ảnh hợp lệ, Enter và Tab có focus','90/90 lượt trang/viewport; luồng mua mobile; dữ liệu tên dài và HTML escape đạt')],
  'sql_note':'Không thêm nghiệp vụ SQL. Test dữ liệu dài thay tạm tên/mô tả sản phẩm trong DB thử, rồi khôi phục nguyên giá trị.'},
 {'id':'CUS-06','title':'Gợi ý hoa theo dịp và ngân sách','points':'2,5','status':'LOCAL_PASS',
  'goal':'Dịp/tone/loại/giá từ DB, URL chia sẻ, empty state gợi ý nới điều kiện, không lộ hidden.',
  'actual':'Giữ biểu mẫu và taxonomy hiện có; thêm kiểm tra ngân sách đảo ngay tại form, lỗi inline và aria-invalid.',
  'files':['includes/customer-recommendations.php','assets/js/customer-forms.js','tests/customer/completion.cjs','docs/evidence/CUS-06/'],
  'pass':'Đạt tổ hợp mẫu, URL reload, hidden và empty state trên mobile; không dùng AI API hoặc sinh sản phẩm giả.',
  'commit':'fix(customer): validate recommendation budgets',
  'image':'CUS-06/report-recommendation-empty.png',
  'checks':[('Kết quả và URL','Chọn Sinh nhật/Hồng/bó hoa/giá tối đa 400000 trả đúng DB','Bộ browser trả Garden; URL tải lại vẫn 1 kết quả'),('Giá đảo và không kết quả','Không gửi khoảng đảo; có cách đặt lại khi không có mẫu','Form chặn khoảng đảo; empty state/URL giữ đúng; sản phẩm hidden không xuất hiện')],
  'sql_note':'Đọc taxonomy từ DB; thử ẩn sản phẩm trong database riêng rồi khôi phục. Chưa thay schema hoặc quy tắc tồn kho.'},
 {'id':'CUS-07','title':'Kiểm thử tài liệu và bàn giao','points':'3','status':'BLOCKED',
  'goal':'Test HTTP/browser/SQL, ảnh thật, báo cáo, PR nhỏ và CI xanh sau review.',
  'actual':'Bổ sung test theo task, SQL tổng hợp, ảnh 18 trang, audit, báo cáo Word, mô tả PR và commit message. Mật khẩu kiểm thử sinh ngẫu nhiên trong bộ nhớ.',
  'files':['tests/customer/http.cjs','tests/customer/browser.cjs','tests/customer/completion.cjs','tests/customer/sql-audit.php','scripts/build-customer-report.py','docs/customer-test-report.md','docs/pr-description.md','docs/BAO_CAO_THANH_VIEN_2_CUSTOMER.docx'],
  'pass':'Test/tài liệu local đã chuẩn bị; không có commit, push, PR, review hoặc CI của nhánh customer vì chưa được người dùng ra lệnh.',
  'commit':'test(customer): add regression evidence and handoff report',
  'image':'CUS-07/report-order-success.png',
  'checks':[('Hồi quy và lint','Các test/lint thực chạy đạt, báo rõ lần thử thất bại','Kết quả cuối lấy từ JSON, kiểm SHA256 3 file core/schema không đổi'),('Bàn giao Git','Chỉ commit/push/PR sau lệnh đúng thao tác','Chưa commit - chờ duyệt; chưa push; chưa mở PR; chưa có CI mới')],
  'sql_note':'9 truy vấn tổng hợp chỉ trả số lượng và sai lệch, không xuất dữ liệu cá nhân. Các đơn TEST-* dùng cho phân trang được loại khỏi kiểm tra tổng tiền đơn CSH-*.'},
]
checklist=[
 ('Mỗi 18 PHP page render HTTP 200 khi được quyền; slug sai/đăng nhập sai xử lý hợp lệ.',True,'HTTP 18 trang; sản phẩm 404; CUS-04 login sai; browser 90 lượt.'),
 ('Filter theo tone, dịp, giá, loại, search nhiều điều kiện chính xác.',True,'CUS-02/results.json: 6 bộ lọc khớp SQL; browser tổ hợp nhiều điều kiện.'),
 ('Guest add/update/delete và refresh giỏ; API Session trả giá trị đồng bộ.',True,'HTTP giỏ; browser queued plus/minus; CUS-03 giỏ sau refresh.'),
 ('Checkout giao hàng và pickup; validation address/date/time/phone; DB tạo đúng orders, order_items.',True,'HTTP checkout; browser mobile delivery; sql-results.json tổng tiền và giao nhận.'),
 ('Xử lý lỗi mạng khi sync giỏ; không hiển thị thành công giả.',True,'Browser abort actions/cart.php, count chưa tăng và có retry.'),
 ('Đặt hàng liên tiếp/double click không gây duplicate không mong muốn.',True,'HTTP hai request cùng key; browser requestSubmit hai lần chỉ thêm 1 đơn.'),
 ('Đăng ký email trùng lỗi rõ; đăng nhập/đăng xuất hợp lệ.',True,'HTTP duplicate; browser signup/profile/logout/login.'),
 ('Customer A chỉ xem được lịch sử A; customer B không truy cập dữ liệu A.',True,'HTTP isolation; success page không lộ đơn A cho B; thử user_id giả.'),
 ('Lookup mã+phone sai bị từ chối mà không leak order.',True,'HTTP 404 khi phone sai; payload không có phone/address; browser lookup.'),
 ('Quên mật khẩu: gửi được email kiểm thử và token có hạn dùng 1 lần.',True,'SMTP local thực; HTTP token hết hạn/replay; CUS-04 reset qua browser.'),
 ('Mobile 375px, tablet 768px, desktop 1440px; không overflow.',True,'18 trang × 5 viewport; ảnh thật ba kích thước trong CUS-01 và từng task.'),
 ('Có test report + ảnh chụp thật + link PR cho mỗi nhóm task.',False,'Có test report và ảnh. Link PR chưa có vì người dùng chưa cho phép mở PR.'),
]
env_text='PHP 8.2.12 của XAMPP; MariaDB cổng 3307; database cas_hoa_customer_preview riêng; PHP development server tại http://127.0.0.1:8033; SMTP local 1025; Chrome do Playwright điều khiển. Dữ liệu test là giả và chỉ dùng nội bộ.'
limits=[
 'URL Apache http://localhost/Lap_trinh_web/index.php không mở được khi kiểm tra; thư mục triển khai htdocs chưa tồn tại. Môi trường preview nằm trong repo, không sửa cấu hình hệ thống. Kết quả local không chứng minh Apache/phpMyAdmin của nhóm đã nghiệm thu.',
 'CUS-02 chưa có gallery nhiều ảnh hoặc tồn kho thành phẩm. Schema chỉ có products.image, chưa có bảng ảnh/quan hệ định mức để suy ra hết hàng. Đã đề xuất lead; không tự đổi schema, role, luật giá hoặc trạng thái đơn.',
 'actions/checkout.php và includes/auth.php đã có diff trước phiên. Test dựa trên bản đó; SHA256 trước/sau bằng nhau. Cần lead review riêng idempotency, snapshot đơn/giấy gói và giới hạn password trước tích hợp.',
 'SMTP local chứng minh gửi/nhận thư thử trong máy; chưa kiểm chứng dịch vụ mail Internet. Không xuất thư, token, mật khẩu hay .env trong báo cáo.',
 'Ảnh là nội dung trình duyệt thật, không có thanh địa chỉ Chrome và không thay thế bộ PDF có thanh địa chỉ nếu giảng viên yêu cầu. Không dựng thanh địa chỉ giả.',
 'Tương phản được đo cụ thể cho nút chính, focus/label/overflow được kiểm tự động và ảnh được xem; chưa audit WCAG toàn diện hoặc mọi công nghệ hỗ trợ.',
 'Chưa commit, push, mở PR hoặc chạy CI nhánh mới. Điểm là tự đánh giá local theo trọng số task, không phải điểm giảng viên và không tương đương tỷ lệ GitHub Contributors.',
]
def count(group): return f"{sum(r['status']=='PASS' for r in group['results'])}/{len(group['results'])}"
summary=f"DONE: 0; LOCAL_PASS: 5; CODE_EXISTS: 0; NOT_DONE: 0; BLOCKED: 2. Tổng trọng số các task đạt local: 19,5/27,5 điểm. Không cộng 5 điểm CUS-02 hoặc 3 điểm CUS-07 khi điều kiện còn thiếu. HTTP/SQL {count(http)}, browser {count(browser)}, PHP lint {count(php)}, JavaScript syntax {count(js)}, SQL tổng hợp {count(sql)}, bổ sung {sum(len(g['results']) for g in supplement.values())} nhóm đều đạt."

# Các tài liệu Markdown lấy cùng dữ liệu với Word để tránh chênh lệch kết luận.
lines=['# Báo cáo kiểm thử Customer','', 'Ngày bắt đầu: 14/10/2026. Nhánh feature/customer-completion.','',env_text,'',summary,'', '## Cách chạy lại','',
 'Chạy tuần tự; không chạy HTTP/browser/SQL có ghi dữ liệu đồng thời. Xem docs/customer-preview.md để đặt CAS_APP_URL, DB_PORT, DB_NAME, PHP_BIN và PLAYWRIGHT_MODULE.','',
 '```text','node tests/customer/http.cjs','node tests/customer/browser.cjs','node tests/customer/completion.cjs CUS-02','node tests/customer/completion.cjs CUS-03','node tests/customer/completion.cjs CUS-04','node tests/customer/completion.cjs CUS-05','node tests/customer/completion.cjs CUS-06','php tests/customer/sql-audit.php','```','',
 'Lưu ý: chỉ database tên cas_hoa_customer_preview hoặc cas_hoa_customer_test_* được helper SQL chấp nhận. Không chạy trên dữ liệu thật. Mật khẩu sinh ngẫu nhiên không ghi ra file.','']
for task in tasks:
    lines.extend([f"## {task['id']} {task['title']}",'',f"Trạng thái: {task['status']}. {task['pass']}",'',task['actual'],'', '| Kiểm tra | Expected | Actual |','|---|---|---|'])
    lines.extend('| '+' | '.join(row)+' |' for row in task['checks'])
    if task['id'] in supplement:
        lines.extend('| '+r['name']+' | '+r['expected']+' | '+r['actual']+' |' for r in supplement[task['id']]['results'])
    lines.extend(['',task['sql_note'],'', 'Tệp: '+', '.join('`'+f+'`' for f in task['files'])+'.','',f"Ảnh: `evidence/{task['image']}`; ảnh khác ở `evidence/{task['id']}/`.",''])
lines.extend(['## SQL kiểm tra và kết quả','',sql['note'],''])
for row in sql['results']:
    expected=str(row.get('expected', '>= '+str(row.get('minimum'))))
    lines.extend([f"### {row['name']}",'','```sql',row['sql']+';','```',f"Expected: {expected}. Actual: {row['actual']}. Kết quả: {row['status']}.",''])
lines.extend(['## Checklist mục 6',''])
lines.extend(f"- [{'x' if done else ' '}] {label} Bằng chứng: {proof}" for label,done,proof in checklist)
lines.extend(['','## Lỗi gặp trong quá trình và xử lý','',
 '- Test reset qua fragment gặp response null khi cùng trang: sửa helper kiểm tra tải lại trang, reset thực qua browser đã đạt.',
 '- Khi thay mật khẩu cố định bằng random, khai báo crypto trùng gây lỗi test HTTP: bỏ khai báo lặp, chạy lại đạt.',
 '- Một lượt tài khoản 768px hết thời gian 7 giây khi chờ networkidle: tăng timeout điều hướng lên 30 giây, giữ nguyên assertion và chạy lại. Kết quả lần thất bại vẫn lưu tại CUS-07/browser-attempt-before-timeout-fix.json.',
 '- Tương phản nút 3,94:1 trước sửa; chỉnh CSS riêng Customer và đo lại 6,32:1.',
 '- Lệnh Git chỉ đọc từng bị hệ thống duyệt tự động chặn do giới hạn sử dụng; sau thời điểm giới hạn đã kiểm tra lại thành công, không đổi git config.',
 '', '## Giới hạn và đề xuất',''])
lines.extend('- '+s for s in limits)
(DOCS/'customer-test-report.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')

pr=['# Mô tả PR dự thảo cho Customer','', 'Chưa mở PR. Nhánh dự kiến: feature/customer-completion → feature/requirement-2-php-mysql.','',
 '## Vấn đề và thay đổi','', 'Hoàn thiện kiểm chứng luồng khách hàng đang có: hiển thị lỗi khoảng giá/ngân sách, nhãn sắp xếp, feedback giỏ/checkout, giữ dữ liệu đăng ký khi có lỗi và tăng tương phản. Thêm audit 18 trang, kiểm thử HTTP/browser/SQL, ảnh thật và báo cáo Word.','',summary,'',
 '## Các phần cần review riêng','', 'Diff checkout/auth đã có trước phiên, không sửa thêm. Gallery và tồn kho thành phẩm cần lead chốt. Xem docs/de-xuat-cho-lead.md. Không gộp mặc nhiên các tệp core vào commit Customer.','',
 '## Chia commit và PR theo task','', '| Task | Commit message đề xuất | Trạng thái commit và PR |','|---|---|---|']
for task in tasks: pr.append(f"| {task['id']} | {task['commit']} | Chưa commit - chờ duyệt; chưa có PR |")
pr.extend(['','Một lệnh duyệt chỉ áp dụng đúng thao tác: commit task CUS-0X, push hoặc mở PR. Không tự staging/commit/push/PR. Các tệp dùng chung như main-r2.js và customer-forms.js cần chia hunk sau khi được phép; không chia commit hình thức để tăng đóng góp.','',
 '## Kiểm chứng','', 'Xem docs/customer-test-report.md, customer-audit.md và docs/evidence/CUS-01..CUS-07. Chưa có CI mới hoặc review của lead. Không viện dẫn CI cũ trên nhánh gốc như CI của thay đổi chưa commit.'])
(DOCS/'pr-description.md').write_text('\n'.join(pr)+'\n',encoding='utf-8')

doc=Document()
section=doc.sections[0];section.page_width=Inches(8.27);section.page_height=Inches(11.69)
section.top_margin=section.bottom_margin=Inches(.65);section.left_margin=section.right_margin=Inches(.68)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(6)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.core_properties.title='Báo cáo kết quả Customer CÁ S HOA'
doc.core_properties.author='Thành viên 2'
def para(text,style=None): return doc.add_paragraph(text,style)
def heading(text,level=1): return doc.add_heading(text,level)
def page(): doc.add_page_break()
def table(headers, rows, widths=None, size=9):
    tbl=doc.add_table(rows=1, cols=len(headers));tbl.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl.autofit=False
    if widths:
        for column,width in zip(tbl.columns,widths):column.width=Inches(width)
    for i,h in enumerate(headers):tbl.rows[0].cells[i].text=h
    for row in rows:
        cells=tbl.add_row().cells
        for i,value in enumerate(row):cells[i].text=str(value)
    for ri,row in enumerate(tbl.rows):
        trpr=row._tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            if widths:cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE8EA' if ri==0 else ('F5F7F7' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
            margin=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margin.append(el)
            tcpr.append(margin)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
                for run in p.runs:run.font.size=Pt(size);run.bold=ri==0
    para('')
    return tbl
def picture(relative,max_height=3.3):
    path=EVIDENCE/relative;assert path.exists(),str(path)
    with Image.open(path) as im:w,h=im.size
    width=min(6.8,max_height*w/h)
    p=para('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Inches(width))
    p=para('Ảnh thật: docs/evidence/'+relative);p.paragraph_format.space_after=Pt(4)
    for run in p.runs:run.font.size=Pt(8)
    # Alt mô tả nội dung, không chứa thông tin riêng.
    inline=doc.inline_shapes[-1]._inline;inline.docPr.set('descr','Ảnh trình duyệt kiểm thử '+relative)

para('CÁ S HOA','Title')
para('Báo cáo kết quả công việc Customer','Title')
para('Thành viên 2','Subtitle')
para('Đề tài CÁ\'S HOA — Website bán hoa tươi')
para('Đối chiếu kế hoạch CUS-01 đến CUS-07')
para('Thời gian thực hiện 14 đến 15/10/2026')
para('Repository https://github.com/tbao4151/Lap_trinh_web')
para('Nhánh làm việc feature/customer-completion')
para('Nhánh tích hợp feature/requirement-2-php-mysql')
para('HEAD khi tiếp tục công việc 9c35a949484c654aa8a7cf9522bcefe09f7ce92c')
heading('Kết luận bàn giao',2)
para('Đã kiểm thử và hoàn thiện các phần Customer có thể xử lý trong phạm vi được giao. Năm task đạt local; hai task còn phụ thuộc lead hoặc lệnh Git của người dùng. Báo cáo trình bày cả phần đạt và phần chưa đủ điều kiện nghiệm thu.')
para('Chưa commit - chờ duyệt. Chưa push và chưa mở PR.')
heading('Cơ sở đối chiếu',2)
para('Kế hoạch 02_PLAN_THANH_VIEN_2_CUSTOMER.md ở thư mục cha workspace; REQUIREMENT2_SETUP.md; tài liệu phạm vi trong repo; PDF NHOM5_52300179_52400271_52400304.pdf gồm 40 trang minh chứng cũ. PDF không phải bằng chứng chạy local trong lần này.')
page();heading('Tóm tắt kết quả')
para(summary)
table(['Trạng thái','Số task','Ý nghĩa'],[['DONE',0,'Chưa có nghiệm thu tích hợp cuối'],['LOCAL_PASS',5,'Thực chạy đạt trên preview của máy này'],['CODE_EXISTS',0,'Không dùng trạng thái có code để thay test'],['NOT_DONE',0,'Các mục đã được xử lý hoặc ghi blocker'],['BLOCKED',2,'CUS-02 schema/contract; CUS-07 PR/CI']], [1.2,.7,4.9])
heading('Môi trường kiểm thử',2);para(env_text)
para('Apache tại URL localhost/Lap_trinh_web chưa chạy. MariaDB mới được khởi tạo trong .local/mysql-completion thuộc repo bằng đường dẫn Windows dạng ngắn ASCII. Không thay database cas_hoa của nhóm.')
table(['Bộ kiểm tra','Đạt','Nguồn'],[['HTTP và SQL nghiệp vụ',count(http),'evidence/customer/http-results.json'],['Browser và responsive',count(browser),'evidence/customer/browser-results.json'],['PHP lint',count(php),'evidence/CUS-07/php-lint.json'],['JavaScript syntax',count(js),'evidence/CUS-07/js-lint.json'],['SQL tổng hợp',count(sql),'evidence/CUS-07/sql-results.json'],['Kiểm tra bổ sung',sum(len(g['results']) for g in supplement.values()),'evidence/CUS-02..CUS-06/results.json']], [2.15,.65,4])
para('Mốc thời gian từng lần chạy có trong JSON. Một lượt timeout browser và lỗi harness đã được ghi nhận, sửa và chạy lại; không xóa dấu vết lần thất bại khỏi phần minh chứng CUS-07.')

matrix_section=doc.add_section(WD_SECTION_START.NEW_PAGE)
matrix_section.orientation=WD_ORIENT.LANDSCAPE;matrix_section.page_width=Inches(11.69);matrix_section.page_height=Inches(8.27)
heading('So sánh kế hoạch với thực tế')
para('Mỗi task được đối chiếu với mục tiêu, trọng số, phần đã xử lý, trạng thái, tệp và điều kiện Pass. Chưa có mã commit mới.')
matrix_goals=['Audit đủ 18 trang và lập bằng chứng.','Catalog đa tiêu chí, chi tiết, ảnh và hết hàng.','Giỏ Session, checkout đúng và chống đơn trùng.','Tài khoản, cách ly A/B và reset email.','5 viewport, bàn phím, nhãn và tương phản.','Gợi ý theo DB, ngân sách và URL chia sẻ.','Test, tài liệu, minh chứng, PR và CI.']
matrix_actual=['18 trang có smoke, ảnh và audit.','Sửa nhãn sắp xếp và lỗi giá; bộ lọc khớp SQL.','Thêm trạng thái bận; giỏ, giá server và retry đạt.','Giữ dữ liệu form lỗi; validation và SMTP/browser đạt.','Đo và sửa tương phản từ 3,94 lên 6,32:1.','Chặn ngân sách đảo; hidden/empty/URL đạt.','Test/lint/SQL đạt; Word và mô tả PR đã soạn.']
matrix_pass=['LOCAL_PASS\nĐạt local; chưa nghiệm thu Apache.','BLOCKED\nCatalog đạt; gallery/tồn kho chờ lead.','LOCAL_PASS\nĐạt local; diff core cũ chờ lead review.','LOCAL_PASS\nĐạt SMTP local; chưa thử mail Internet.','LOCAL_PASS\nĐạt phạm vi test; chưa audit WCAG toàn diện.','LOCAL_PASS\nĐạt các tổ hợp đã test trên DB thật.','BLOCKED\nLocal đạt; commit/PR/CI chưa được phép.']
table(['Task và điểm','Mục tiêu kế hoạch','Thực tế','Trạng thái và Pass','Tệp mới hoặc sửa trong phiên','Commit'], [[t['id']+'\n'+t['points']+' điểm',matrix_goals[i],matrix_actual[i],matrix_pass[i],'\n'.join(t['files']),'Chưa commit - chờ duyệt'] for i,t in enumerate(tasks)], [.65,1.5,2.1,1.9,2.6,1.0],size=8.5)
body_section=doc.add_section(WD_SECTION_START.NEW_PAGE)
for p in doc.paragraphs[-2:]:
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=Pt(1)
    p.paragraph_format.keep_with_next=False
    p.paragraph_format.page_break_before=False
body_section.orientation=WD_ORIENT.PORTRAIT;body_section.page_width=Inches(8.27);body_section.page_height=Inches(11.69)
for task_index,task in enumerate(tasks):
    if task_index:page()
    heading(task['id']+' '+task['title'])
    para('Trọng số '+task['points']+' điểm. Trạng thái '+task['status']+'.')
    para(task['actual']);para(task['pass'])
    table(['Kiểm tra','Expected','Actual'],task['checks'],[1.25,2.55,3],size=9)
    para(task['sql_note'])
    para('Tệp mới/sửa trong lần tiếp tục: '+', '.join(task['files'])+'.')
    para('Commit: Chưa commit - chờ duyệt. Đề xuất: '+task['commit']+'.')
    picture(task['image'],3.6)

page();heading('Audit 18 trang khách hàng')
para('Nguồn: docs/customer-audit.md. Mỗi trang có HTTP 200 khi được quyền và 5 lượt kiểm tra responsive; tài khoản/lịch sử được mở trong phiên đăng nhập. Các trang lỗi được kiểm tra riêng theo use case.')
audit_rows=[]
for entry in http['results']:
    if entry['name'].startswith('HTTP page '):
        route=entry['name'][10:];name=route.split('?')[0]
        checks=[r for r in browser['results'] if r['name'].startswith('Responsive ') and r['name'].endswith('px '+route)]
        assert len(checks)==5 and all(r['status']=='PASS' for r in checks)
        audit_rows.append([name,'LOCAL_PASS','HTTP 200; 5 viewport; ảnh 375/768/1440'])
table(['Trang PHP','Trạng thái','Bằng chứng'],audit_rows,[2.4,1.15,3.25],size=9)
para('Ảnh trong evidence/CUS-01 theo tên trang và chiều rộng. Riêng liên hệ chụp biểu mẫu thật để tránh đưa thông tin liên hệ cửa hàng vào tài liệu.')

page();heading('Checklist kiểm thử bắt buộc')
para('Đối chiếu nguyên phạm vi mục 6 của kế hoạch. Dấu đạt chỉ áp dụng môi trường local đã mô tả.')
table(['Đạt','Yêu cầu','Bằng chứng'],[['Có' if done else 'Chưa',label,proof] for label,done,proof in checklist],[.5,3.3,3],size=9)

page();heading('SQL kiểm tra và kết quả')
para(sql['note'])
for row in sql['results']:
    heading(row['name'],2)
    p=para(row['sql']+';')
    for run in p.runs:run.font.name='Consolas';run.font.size=Pt(8.5)
    expected=str(row.get('expected','>= '+str(row.get('minimum'))))
    para(f"Expected {expected}; actual {row['actual']}; {row['status']}.")

page();heading('Phần còn tồn đọng và đề xuất lead')
for item in limits:para(item, 'List Bullet')
heading('Lỗi đã xử lý',2)
para('Đã sửa khoảng giá/ngân sách đảo không có thông báo rõ, nhãn sắp xếp chưa chính xác, và tương phản nút chưa đạt. Các lỗi test do điều hướng fragment, biến crypto trùng và timeout 7 giây được sửa trong harness, sau đó kiểm thử lại. Không coi việc sửa harness là sửa lỗi nghiệp vụ ứng dụng.')
para('Chi tiết đề xuất: docs/de-xuat-cho-lead.md. Không phát hành thay đổi schema/core để tự tháo blocker.')

page();heading('Commit PR và mức đóng góp')
table(['Task','Tệp mới hoặc sửa trong phiên','Commit và PR'], [[t['id'],', '.join(t['files']),'Chưa commit - chờ duyệt\nChưa có PR'] for t in tasks],[.7,4.55,1.55],size=8.5)
para('Mô tả PR và commit message theo task được chuẩn bị tại docs/pr-description.md. Chưa có mã commit mới, URL PR, review hoặc CI mới để dẫn chứng.')
heading('Đánh giá đóng góp',2)
para('19,5/27,5 điểm là tổng trọng số các task đủ bằng chứng local, không phải số điểm đã được chấm. Repo đã có nhiều code chưa commit trước phiên này; không nhận toàn bộ diff đó là code mới viết trong lần tiếp tục. Đóng góp mới tập trung vào sửa UX có lỗi cụ thể, tăng tương phản, mở rộng kiểm thử và chuẩn hóa minh chứng/tài liệu.')
para('Trước khi nộp cần lead duyệt gallery/tồn kho và diff core, triển khai kiểm tra Apache đúng môi trường nhóm, rồi người dùng ra lệnh riêng cho commit, push, mở PR. Không cam kết tỷ lệ Contributors dựa trên số file hoặc số commit.')
doc.save(DOCS/'BAO_CAO_THANH_VIEN_2_CUSTOMER.docx')

# Kiểm tra cấu trúc ngay sau ghi; kiểm tra hình thức thực hiện bằng bản render riêng.
reopened=Document(DOCS/'BAO_CAO_THANH_VIEN_2_CUSTOMER.docx')
assert len(reopened.inline_shapes)==7
assert len(audit_rows)==18
assert len(reopened.tables)>=12
(EVIDENCE/'CUS-07/docx-structure.json').write_text(json.dumps({'tables':len(reopened.tables),'images':len(reopened.inline_shapes),'audit_pages':len(audit_rows),'status':'PASS','note':'Đã mở lại bằng python-docx; hình thức cần kiểm tra bản render.'},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(DOCS/'BAO_CAO_THANH_VIEN_2_CUSTOMER.docx'),'tables':len(reopened.tables),'images':len(reopened.inline_shapes)},ensure_ascii=False))
